"""使用临时台账和资料验证逐条归档流程。"""

from __future__ import annotations

import hashlib
import json
import os
import sys
import tempfile
import threading
import time
import unittest
import urllib.request
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader
from reportlab.pdfgen import canvas
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from b24_design_archiver.bluebeam_cli import combine_and_print_a4, find_script_engine
from b24_design_archiver.config import load_config
from b24_design_archiver.operations import _copy_file, filesystem_path
from b24_design_archiver.records import specialty_for
from b24_design_archiver.server import ArchiveState, create_server


class WorkflowTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        workbook = self.root / "台账.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append([
            "审批表编号", "提出单位一", "变更原因二", "变更内容三", "变更通知单编号",
            "变更编号和说明", "日期五", "六版本", "是否完成面单", "何种专业四",
        ])
        sheet.append(["BG001", "设计部", "图纸完善", "调整水泵参数", "05-03-C2-001", "05-03-C2-001（水泵）", "2025-09-30", "施工图", "是", "给排水"])
        sheet.append(["BG002", "设计部", "图纸完善", "修改管线", "05-03-C2-002", "05-03-C2-002（管线）", "2025-10-02", "施工图", "否", "给排水"])
        sheet.append([None] * 10)
        book.save(workbook)
        before = self.root / "前.docx"
        after = self.root / "后.docx"
        for path in (before, after):
            doc = Document()
            doc.add_paragraph("变更通知单编号 日期五")
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "变更内容三"
            doc.save(path)
        source = self.root / "source"
        target = source / "05-03-C2-001（水泵）"
        target.mkdir(parents=True)
        (target / "05-03-C2-001说明.docx").write_bytes(b"sample docx")
        pdf = target / "05-03-C2-001图纸.pdf"
        draw = canvas.Canvas(str(pdf), pagesize=(900, 500))
        draw.drawString(30, 30, "drawing")
        draw.save()
        (source / "05-03-C2-0010误匹配.pdf").write_bytes(b"ignored")
        self.config = {
            "workbook": workbook, "template_before": before, "template_after": after,
            "template_change_date": "2025-10-01", "output_root": self.root / "output",
            "source_directories": [source],
        }
        with patch("b24_design_archiver.server.load_config", return_value=self.config), patch("b24_design_archiver.server.SETTINGS_FILE", self.root / "settings.json"):
            self.state = ArchiveState()
        self.first_id = self.state.records[0]["_id"]
        self.second_id = self.state.records[1]["_id"]

    def test_record_preview_and_actions(self) -> None:
        self.assertEqual(len(self.state.records), 2)
        preview = self.state.preview(self.first_id)
        self.assertEqual(len(preview["directories"]), 1)
        self.assertEqual(len(preview["pdfs"]), 1)
        self.assertIn("水泵", preview["output_folder"])
        result = self.state.action(self.first_id, "generate")
        generated = Path(result["files"][0])
        doc = Document(generated)
        self.assertIn("调整水泵参数", doc.tables[0].cell(0, 0).text)
        self.assertIn("2025.09.30", doc.paragraphs[0].text)
        copy_result = self.state.action(self.first_id, "copy")
        self.assertEqual(len(copy_result["files"]), 3)
        self.assertTrue(Path(copy_result["files"][0]).is_dir())
        self.assertEqual(self.state.action(self.first_id, "copy")["files"], [])
        archive_result = self.state.action(self.first_id, "zip")
        with zipfile.ZipFile(archive_result["files"][0]) as archive:
            self.assertEqual(len(archive.namelist()), 2)
        self.assertEqual(self.state.action(self.first_id, "zip")["files"], archive_result["files"])
        second = self.state.action(self.second_id, "generate")
        self.assertNotEqual(Path(second["files"][0]).parent, generated.parent)

    def test_records_json_follows_output_and_reload(self) -> None:
        first = self.state.records_json
        self.assertEqual(first, self.config["output_root"] / "B24变更条目.json")
        exported = json.loads(first.read_text(encoding="utf-8"))
        self.assertEqual(len(exported), len(self.state.records))
        self.assertEqual(exported[0]["_id"], hashlib.sha256(b"05-03-C2-001").hexdigest())
        self.assertEqual(exported[0]["专业"], "给排水")
        self.assertEqual(exported[0]["变更内容"], "调整水泵参数")
        self.assertEqual(exported[0]["版本"], "施工图")
        for marker in ("审批表编号", "提出单位一", "变更原因二", "变更内容三", "何种专业四", "日期五", "六版本"):
            self.assertNotIn(marker, exported[0])
        new_root = self.root / "new-output"
        with patch("b24_design_archiver.server.SETTINGS_FILE", self.root / "settings.json"):
            self.state.set_output_root(str(new_root))
        self.assertEqual(self.state.records_json, new_root / "B24变更条目.json")
        self.assertEqual(json.loads(self.state.records_json.read_text(encoding="utf-8")), exported)
        self.state.records_json.unlink()
        self.state.reload()
        self.assertTrue(self.state.records_json.is_file())
        self.assertEqual(self.state.records[0]["_id"], self.first_id)

    def test_specialty_from_change_number(self) -> None:
        cases = {
            "05-03-C2-001": "给排水", "06-00-C2-001": "暖通",
            "JZ-05-00-C2-001": "精装给排水", "JZ-06-00-C2-001": "精装暖通",
            "R20-03-Z-01": "热力", "RL20-03-Z-03": "热力",
            "251614GC-01GS-1": "燃气", "泳池变更-01": "泳池设备",
        }
        for code, expected in cases.items():
            self.assertEqual(specialty_for(code), expected)

    def test_sync_materials_updates_json(self) -> None:
        self.state.start_material_sync()
        deadline = time.monotonic() + 15
        while self.state.sync_status["running"] and time.monotonic() < deadline:
            time.sleep(0.05)
        self.assertFalse(self.state.sync_status["running"])
        self.assertEqual(self.state.sync_status["completed"], 2)
        self.assertEqual(self.state.sync_status["errors"], [])
        exported = json.loads(self.state.records_json.read_text(encoding="utf-8"))
        self.assertTrue(exported[0]["已复制资料"])
        self.assertTrue(any(path.endswith("图纸.pdf") for path in exported[0]["资料文件"]))
        self.assertFalse(exported[1]["已复制资料"])
        self.assertEqual(self.state.start_material_sync()["running"], True)
        while self.state.sync_status["running"] and time.monotonic() < deadline:
            time.sleep(0.05)
        self.assertEqual(self.state.sync_status["errors"], [])

    def test_legacy_config_directories_are_combined(self) -> None:
        extra = self.root / "legacy-docx"
        extra.mkdir()
        legacy = self.root / "legacy.yaml"
        legacy.write_text(yaml.safe_dump({"base_directories": [str(self.config["source_directories"][0])], "docx_directories": [str(extra)]}), encoding="utf-8")
        settings = self.root / "app.yaml"
        app = {**self.config, "source_directories": [str(path) for path in self.config["source_directories"]], "legacy_config_files": [str(legacy)]}
        for key in ("workbook", "template_before", "template_after", "output_root"):
            app[key] = str(app[key])
        settings.write_text(yaml.safe_dump({"app": app}, allow_unicode=True), encoding="utf-8")
        with patch("b24_design_archiver.config.CONFIG_PATH", settings):
            loaded = load_config()
        self.assertEqual(loaded["source_directories"], [self.config["source_directories"][0], extra])

    @unittest.skipUnless(os.name == "nt", "Windows 长路径行为")
    def test_copy_file_with_long_output_path(self) -> None:
        source = self.root / "source.txt"
        source.write_text("long path", encoding="utf-8")
        parent = self.root / ("x" * 180)
        target = parent / ("y" * 80 + ".txt")
        self.assertGreater(len(str(target)), 260)
        self.addCleanup(parent.rmdir)
        self.addCleanup(filesystem_path(target).unlink, missing_ok=True)
        _copy_file(source, target)
        self.assertEqual(filesystem_path(target).read_text(encoding="utf-8"), "long path")

    def test_http(self) -> None:
        server = create_server(self.state)
        worker = threading.Thread(target=server.serve_forever, daemon=True)
        worker.start()
        try:
            base = f"http://127.0.0.1:{server.server_port}"
            with urllib.request.urlopen(base + "/api/records") as response:
                data = json.load(response)
            self.assertEqual(len(data["records"]), 2)
            self.assertEqual(data["records_json"], str(self.state.records_json))
            request = urllib.request.Request(base + "/api/action", data=json.dumps({"id": self.first_id, "action": "generate"}).encode(), headers={"Content-Type": "application/json", "X-Session-Token": self.state.token})
            with urllib.request.urlopen(request) as response:
                result = json.load(response)
            self.assertEqual(len(result["files"]), 1)
            unopened = self.state.folder(self.state.get_record(self.second_id))
            self.assertFalse(unopened.exists())
            opener = "b24_design_archiver.server.os.startfile" if os.name == "nt" else "b24_design_archiver.server.subprocess.Popen"
            with patch(opener) as open_folder:
                request = urllib.request.Request(base + "/api/open", data=json.dumps({"id": self.second_id}).encode(), headers={"Content-Type": "application/json", "X-Session-Token": self.state.token})
                with urllib.request.urlopen(request) as response:
                    opened = json.load(response)
                open_folder.assert_called_once()
            self.assertEqual(opened["folder"], str(unopened))
            self.assertTrue(unopened.is_dir())
        finally:
            server.shutdown()
            server.server_close()

    @unittest.skipUnless(os.name == "nt", "Bluebeam CLI 仅支持 Windows")
    def test_bluebeam_cli_merge_and_a4(self) -> None:
        try:
            find_script_engine()
        except FileNotFoundError as exc:
            self.skipTest(str(exc))
        source = Path(self.state.preview(self.first_id)["pdfs"][0])
        approval = self.root / "approval.pdf"
        draw = canvas.Canvas(str(approval), pagesize=(595, 842))
        draw.drawString(30, 30, "approval")
        draw.save()
        merged = self.root / "merged.pdf"
        a4 = self.root / "a4.pdf"
        combine_and_print_a4([approval, source], merged, a4)
        merged_pages = PdfReader(merged).pages
        self.assertEqual(len(merged_pages), 2)
        self.assertIn("approval", merged_pages[0].extract_text())
        pages = PdfReader(a4).pages
        self.assertEqual(len(pages), 2)
        for page in pages:
            self.assertAlmostEqual(float(page.mediabox.width), 595.3, delta=2)
            self.assertAlmostEqual(float(page.mediabox.height), 841.9, delta=2)

    def test_multiple_pdfs_require_selection_before_writing(self) -> None:
        folder = self.root / "source" / "05-03-C2-001（水泵）"
        (folder / "05-03-C2-001附件.pdf").write_bytes(b"second candidate")
        self.state.sources.invalidate()
        with self.assertRaisesRegex(ValueError, "选择要合并的文件"):
            self.state.action(self.first_id, "merge")
        self.assertFalse(self.state.folder(self.state.get_record(self.first_id)).exists())
        with self.assertRaisesRegex(ValueError, "请选择当前变更"):
            self.state.action(self.first_id, "merge", ["outside.pdf"])


if __name__ == "__main__":
    unittest.main()
