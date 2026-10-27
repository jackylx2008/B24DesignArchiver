"""单条变更的审批表、资料复制、压缩与 PDF 合并。"""

from __future__ import annotations

import hashlib
import os
import re
import shutil
import tempfile
import time
import zipfile
from datetime import date
from pathlib import Path

from docx import Document

from .bluebeam_cli import combine_and_print_a4


def safe_name(value: str, limit: int = 96) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    return value[:limit].rstrip(" .") or "未命名"


def filesystem_path(path: Path) -> Path:
    """让 Windows 文件 API 能处理超过传统 MAX_PATH 的成果文件。"""
    if os.name == "nt" and len(str(path)) >= 248 and not str(path).startswith("\\\\?\\"):
        return Path("\\\\?\\" + str(path.resolve()))
    return path


def record_directory(record: dict, output_root: Path) -> Path:
    raw_id = record["变更通知单编号"]
    change_id = safe_name(raw_id, 48)
    description = record.get("变更编号和说明", "")
    if description.startswith(raw_id):
        description = description[len(raw_id):].strip("（）() _-")
    return output_root / safe_name(f"{change_id}_{description}" if description else change_id)


def _template(config: dict, record: dict) -> Path:
    raw_date = record["日期五"]
    if not raw_date:
        raise ValueError("台账日期为空，无法选择审批表模板")
    chosen = "template_after" if date.fromisoformat(raw_date[:10]) >= date.fromisoformat(config["template_change_date"]) else "template_before"
    template = config[chosen]
    if not template.is_file():
        raise FileNotFoundError(f"找不到模板：{template}")
    return template


def _replace_runs(paragraph, values: dict[str, str]) -> None:
    for run in paragraph.runs:
        content = run.text
        for key, replacement in values.items():
            content = content.replace(key, replacement)
        run.text = content


def generate_approval(record: dict, config: dict, folder: Path) -> Path:
    template = _template(config, record)
    target = folder / f"{safe_name(record['变更通知单编号'])}_设计变更审批表.docx"
    if target.exists():
        return target
    values = {key: value for key, value in record.items() if not key.startswith("_")}
    values["日期五"] = values["日期五"].replace("-", ".")
    values["#####"] = record["变更通知单编号"]
    document = Document(template)
    for paragraph in document.paragraphs:
        _replace_runs(paragraph, values)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_runs(paragraph, values)
    folder.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=folder, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        document.save(temporary)
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    return target


def _same_file(first: Path, second: Path) -> bool:
    if filesystem_path(first).stat().st_size != filesystem_path(second).stat().st_size:
        return False

    def digest(path: Path) -> bytes:
        hasher = hashlib.sha256()
        with filesystem_path(path).open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                hasher.update(chunk)
        return hasher.digest()
    return digest(first) == digest(second)


def _unique_target(target: Path, source: Path) -> Path | None:
    if not filesystem_path(target).exists():
        return target
    if _same_file(source, target):
        return None
    for number in range(2, 1000):
        option = target.with_name(f"{target.stem}_{number}{target.suffix}")
        if not filesystem_path(option).exists():
            return option
        if _same_file(source, option):
            return None
    raise ValueError(f"同名文件过多：{target.name}")


def _copy_file(source: Path, target: Path) -> None:
    """同目录暂存后发布，避免同步软件扫描时留下半写入文件。"""
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=target.parent, suffix=target.suffix, delete=False) as handle:
        staged = Path(handle.name)
    try:
        shutil.copy2(source, staged)
        for attempt in range(5):
            try:
                filesystem_path(staged).replace(filesystem_path(target))
                return
            except PermissionError:
                if attempt == 4:
                    raise
                time.sleep(0.2 * (attempt + 1))
    finally:
        staged.unlink(missing_ok=True)


def copy_materials(preview: dict, folder: Path) -> list[str]:
    target_dir = folder / "资料文件"
    copied = []
    for name in preview["directories"]:
        source = Path(name)
        source_hash = hashlib.sha1(str(source.resolve()).encode("utf-8")).hexdigest()[:8]
        target = folder / "原始资料目录" / f"{safe_name(source.name)}_{source_hash}"
        new_files = 0
        for path in source.rglob("*"):
            if path.is_dir():
                (target / path.relative_to(source)).mkdir(parents=True, exist_ok=True)
            elif path.is_file():
                destination = _unique_target(target / path.relative_to(source), path)
                if destination is not None:
                    _copy_file(path, destination)
                    new_files += 1
        if new_files:
            copied.append(str(target))
    sources = sorted(set(preview["pdfs"] + preview["docxs"]))
    for name in sources:
        source = Path(name)
        target_dir.mkdir(parents=True, exist_ok=True)
        target = _unique_target(target_dir / safe_name(source.name, 180), source)
        if target is not None:
            _copy_file(source, target)
            copied.append(str(target))
    return copied


def zip_directory(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".zip", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for file in sorted(source.rglob("*")):
                if file.is_file():
                    archive.write(file, file.relative_to(source))
        with zipfile.ZipFile(temporary) as archive:
            if archive.testzip() is not None:
                raise ValueError("生成的 ZIP 校验失败")
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


def zip_materials(preview: dict, folder: Path) -> list[str]:
    results = []
    for source_name in preview["directories"]:
        source = Path(source_name)
        source_hash = hashlib.sha1(str(source.resolve()).encode("utf-8")).hexdigest()[:8]
        target = folder / "资料压缩包" / f"{safe_name(source.name)}_{source_hash}.zip"
        if target.exists():
            results.append(str(target))
            continue
        zip_directory(source, target)
        results.append(str(target))
    return results


def _word_to_pdf(source: Path, destination: Path) -> None:
    if os.name != "nt":
        raise RuntimeError("Word 转 PDF 目前需要 Windows 和桌面版 Microsoft Word")
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(str(source.resolve()), ReadOnly=True)
        document.ExportAsFixedFormat(str(destination.resolve()), 17)
    finally:
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


def merge_pdfs(approval_docx: Path, preview: dict, folder: Path) -> list[str]:
    pdfs = [Path(name) for name in preview["pdfs"]]
    if not pdfs:
        raise ValueError("未找到对应的原始 PDF，无法合并")
    output = folder / f"{safe_name(approval_docx.stem)}_合并.pdf"
    a4 = folder / f"{safe_name(approval_docx.stem)}_合并_A4.pdf"
    if output.exists() or a4.exists():
        raise FileExistsError("合并 PDF 已存在；为避免覆盖，请先在输出目录中处理旧成果")
    with tempfile.TemporaryDirectory(prefix="b24-word-") as temporary_dir:
        temporary = Path(temporary_dir)
        approval_pdf = temporary / "approval.pdf"
        _word_to_pdf(approval_docx, approval_pdf)
        merged = temporary / "merged.pdf"
        normalized = temporary / "a4.pdf"
        combine_and_print_a4([approval_pdf, *pdfs], merged, normalized)
        staged_outputs: list[tuple[Path, Path]] = []
        published: list[Path] = []
        try:
            for source, destination in ((merged, output), (normalized, a4)):
                with tempfile.NamedTemporaryFile(dir=folder, suffix=".pdf", delete=False) as handle:
                    staged = Path(handle.name)
                staged_outputs.append((staged, destination))
                shutil.copy2(source, staged)
            for staged, destination in staged_outputs:
                staged.rename(destination)
                published.append(destination)
        except OSError:
            for destination in published:
                destination.unlink(missing_ok=True)
            raise
        finally:
            for staged, _ in staged_outputs:
                staged.unlink(missing_ok=True)
    return [str(output), str(a4)]


def zip_selected_folders(folders: list[Path], destination: Path) -> Path:
    if destination.exists():
        raise FileExistsError(f"总 ZIP 已存在：{destination}")
    if not any(folder.is_dir() for folder in folders):
        raise ValueError("所选变更尚无成果目录")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for folder in folders:
            if not folder.is_dir():
                continue
            for file in sorted(folder.rglob("*")):
                if file.is_file():
                    archive.write(file, Path(folder.name) / file.relative_to(folder))
    return destination
