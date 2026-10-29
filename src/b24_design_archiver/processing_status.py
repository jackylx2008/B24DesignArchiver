"""从本地多工作表台账提取变更处理情况。"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .records import _write_json

PARTIES = ("总包单位", "施工单位", "监理单位", "建设单位")
DEFAULT_PARTIES = ("总包单位", "监理单位", "建设单位")


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_processing_status(workbook: Path, change_numbers: set[str]) -> dict[str, list[dict]]:
    """精确匹配 B 列编号，保留各 sheet 的原始处理标记。"""
    wanted = {number.strip().upper() for number in change_numbers}
    results: dict[str, list[dict]] = defaultdict(list)
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        for sheet in book:
            headers = [_text(value) for value in next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))]
            parties = {name: index for index, name in enumerate(headers) if name in PARTIES}
            approval = next((index for index, name in enumerate(headers) if name.startswith("业主是否审批")), None)
            if not parties and approval is None:
                continue
            for row_number, cells in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
                number = _text(cells[1]).upper() if len(cells) > 1 else ""
                if number not in wanted:
                    continue
                status: dict = {"来源文件": workbook.name, "来源工作表": sheet.title, "行号": row_number}
                if approval is not None:
                    status["业主是否审批"] = _text(cells[approval]) if approval < len(cells) else ""
                status["蓝图下发情况"] = {
                    name: _text(cells[index]) if index < len(cells) else ""
                    for name, index in parties.items()
                }
                results[number].append(status)
        return dict(results)
    finally:
        book.close()


def apply_processing_status(workbook: Path, output_root: Path, selected_numbers: set[str]) -> dict:
    """更新根目录总表里的处理情况，不触碰子目录 record.json。"""
    destination = output_root / "B24变更条目.json"
    import json
    records = json.loads(destination.read_text(encoding="utf-8"))
    by_number = {item["变更通知单编号"].upper(): item for item in records}
    unknown = selected_numbers - set(by_number)
    if unknown:
        raise ValueError(f"总表没有这些编号：{', '.join(sorted(unknown))}")
    matched = read_processing_status(workbook, selected_numbers)
    for number in selected_numbers:
        by_number[number]["处理情况"] = summarize_processing_status(
            matched.get(number, []), is_jz=number.startswith("JZ-")
        )
    _write_json(records, destination)
    return {"selected": len(selected_numbers), "matched": sum(bool(matched.get(number)) for number in selected_numbers)}


def summarize_processing_status(rows: list[dict], *, is_jz: bool = False) -> dict:
    """三家单位默认未下发；Excel 可见的勾选覆盖默认值，并保留逐表来源。"""
    parties = {name: "×" for name in DEFAULT_PARTIES}
    specialty_sources = []
    approvals = []
    for row in rows:
        if "业主是否审批" in row:
            value = row["业主是否审批"]
            if value and value not in approvals:
                approvals.append(value)
        if is_jz and "精装" in row.get("来源工作表", "") and row.get("来源工作表") not in specialty_sources:
            specialty_sources.append(row["来源工作表"])
        for name, value in row.get("蓝图下发情况", {}).items():
            if name not in parties:
                parties[name] = "×"
            if value in {"√", "✓", "✔", "☑"}:
                parties[name] = "√"

    result = {"蓝图下发情况": parties, "Excel来源记录": rows}
    if approvals:
        result["业主是否审批"] = approvals
    if specialty_sources:
        result["精装表标记"] = specialty_sources
    return result


def default_processing_status() -> dict:
    return summarize_processing_status([])
